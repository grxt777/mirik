"""Чтение и создание документов для агентов Мирика.

read  <path>                 -> текст PDF / DOCX / PPTX / XLSX / TXT в stdout
docx  <path>  (markdown в stdin) -> собрать .docx из простого markdown
"""
import re
import sys
from pathlib import Path


def read_pdf(p):
    from pypdf import PdfReader
    reader = PdfReader(p)
    out = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        out.append(f"--- стр. {i} ---\n{text}")
    body = "\n\n".join(out)
    if len(re.sub(r"--- стр\. \d+ ---|\s", "", body)) < 20:
        body += "\n\n[В PDF почти нет текста — похоже на скан/картинки. Текст отсюда не извлечь.]"
    return body


def read_docx(p):
    import docx
    d = docx.Document(p)
    out = []
    for para in d.paragraphs:
        if para.text.strip():
            style = (para.style.name or "").lower()
            out.append(("# " if style.startswith("heading") or style.startswith("title") else "") + para.text)
    for t, table in enumerate(d.tables, 1):
        out.append(f"\n[Таблица {t}]")
        for row in table.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def read_pptx(p):
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(p).slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        out.append(f"--- слайд {i} ---\n" + "\n".join(texts))
    return "\n\n".join(out)


def read_xlsx(p):
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"--- лист {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def read(path):
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        return read_pdf(p)
    if ext == ".docx":
        return read_docx(p)
    if ext == ".pptx":
        return read_pptx(p)
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(p)
    if ext in (".doc", ".ppt", ".xls", ".hwp"):
        return f"[Старый формат {ext} не поддерживается — нужен .docx/.pptx/.xlsx или PDF]"
    return p.read_text(errors="replace")


def inline(par, text):
    """**жирный**, *курсив*, `код` внутри абзаца."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def cells(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(d, rows):
    rows = [r for r in rows if not re.fullmatch(r"[\s|:\-]+", r)]  # строка-разделитель |---|
    data = [cells(r) for r in rows]
    width = max(len(r) for r in data)
    t = d.add_table(rows=len(data), cols=width)
    t.style = "Table Grid"
    for i, row in enumerate(data):
        for j in range(width):
            inline(t.cell(i, j).paragraphs[0], row[j] if j < len(row) else "")
            if i == 0:
                for run in t.cell(i, j).paragraphs[0].runs:
                    run.bold = True
    d.add_paragraph()


def make_docx(path, md):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    d.styles["Normal"].font.name = "Times New Roman"
    d.styles["Normal"].font.size = Pt(12)
    code, table = None, []
    for line in md.splitlines():
        if line.strip().startswith("|") and code is None:
            table.append(line)
            continue
        if table:
            add_table(d, table)
            table = []
        if re.fullmatch(r"\s*([-*_])\1{2,}\s*", line):  # --- горизонтальная черта
            continue
        if line.strip().startswith("```"):
            if code is None:
                code = []
            else:
                p = d.add_paragraph()
                r = p.add_run("\n".join(code))
                r.font.name = "Consolas"
                r.font.size = Pt(10)
                code = None
            continue
        if code is not None:
            code.append(line)
            continue
        h = re.match(r"^(#{1,4})\s+(.*)", line)
        if h:
            d.add_heading(h.group(2).strip(), level=min(len(h.group(1)), 3))
        elif re.match(r"^\s*[-*•]\s+", line):
            inline(d.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•]\s+", "", line))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            inline(d.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line))
        elif line.strip():
            inline(d.add_paragraph(), line)
    if table:
        add_table(d, table)
    if code:
        d.add_paragraph("\n".join(code))
    d.save(path)
    return f"Готово: {path}"


if __name__ == "__main__":
    cmd, path = sys.argv[1], sys.argv[2]
    try:
        print(read(path) if cmd == "read" else make_docx(path, sys.stdin.read()))
    except ImportError as e:
        print(f"[Не установлена библиотека: {e.name}]")
        sys.exit(1)
